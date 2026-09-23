# -*- coding: utf-8 -*-
import io
import datetime
from odoo import http
from odoo.http import request, content_disposition


class GanttExportController(http.Controller):

    @http.route('/custom_gantt/export_excel/<int:project_id>',
                type='http', auth='user', methods=['GET'])
    def export_excel(self, project_id, **kwargs):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import (PatternFill, GradientFill, Font, Alignment,
                                          Border, Side)
            from openpyxl.styles.fills import GradientFill, Stop
            from openpyxl.utils import get_column_letter
            from openpyxl.drawing.image import Image as XLImage
        except ImportError:
            return request.make_response(
                "openpyxl not installed. Run: pip install openpyxl",
                headers=[('Content-Type', 'text/plain')]
            )

        project = request.env['gantt.project'].sudo().browse(project_id)
        if not project.exists():
            return request.make_response("Project not found", headers=[('Content-Type', 'text/plain')])

        # Invalidate ORM cache so we always get fresh data on each export
        request.env['gantt.task'].invalidate_model()
        request.env['apqp.timeline.format'].invalidate_model()

        all_tasks = request.env['gantt.task'].sudo().with_context(
            no_recompute=False
        ).search(
            [('project_id', '=', project_id)],
            order='id asc'
        )

        excluded_task_ids = request.env['apqp.timeline.format'].sudo().search([
            ('include_in_report', '=', False)
        ]).mapped('gantt_task_id').ids

        tasks = all_tasks.filtered(lambda t: t.id not in excluded_task_ids)

        # ── De-duplicate legacy phase roots ─────────────────────────────
        # Some projects end up with two root tasks sharing the same name
        # (e.g. "Phase 1: Plan and Define Program"): an old placeholder
        # created before this project used APQP sync (no
        # apqp_timeline_format_id, no attachments), and the real one
        # created/synced by apqp.timeline.chart.action_create_sync_gantt()
        # (linked via apqp_timeline_format_id, carries documents/
        # attachments). When both exist for the same phase name, keep only
        # the APQP-linked one and drop the placeholder (and its children).
        linked_task_ids = set(
            request.env['apqp.timeline.format'].sudo().search(
                [('gantt_task_id', '!=', False)]
            ).mapped('gantt_task_id').ids
        )

        roots_all = [t for t in tasks if not t.parent_id]
        by_name = {}
        for r in roots_all:
            by_name.setdefault(r.name, []).append(r)

        drop_ids = set()
        for name, group in by_name.items():
            if len(group) > 1:
                linked = [t for t in group if t.id in linked_task_ids]
                # Only drop siblings if we can clearly identify which one
                # is the correct (APQP-linked) copy.
                if linked and len(linked) < len(group):
                    keep_ids = {t.id for t in linked}
                    drop_ids.update(t.id for t in group if t.id not in keep_ids)

        if drop_ids:
            def _collect_descendants(task_id, acc):
                for t in tasks:
                    if t.parent_id and t.parent_id.id == task_id:
                        acc.add(t.id)
                        _collect_descendants(t.id, acc)

            for tid in list(drop_ids):
                _collect_descendants(tid, drop_ids)

            tasks = tasks.filtered(lambda t: t.id not in drop_ids)

        # ── Date range ────────────────────────────────────────────────
        all_starts = [t.date_start for t in tasks if t.date_start]
        all_stops  = [t.date_stop  for t in tasks if t.date_stop]
        if not all_starts:
            return request.make_response("No tasks with dates found",
                                          headers=[('Content-Type', 'text/plain')])

        range_start = min(all_starts).date()
        range_end   = max(all_stops).date()
        total_days  = (range_end - range_start).days + 1

        # ── Build WBS tree ────────────────────────────────────────────
        rows = []
        def walk(task, level, parent_serial):
            children = [t for t in tasks if t.parent_id and t.parent_id.id == task.id]
            if parent_serial:
                siblings = [t for t in tasks
                            if t.parent_id and t.parent_id.id == task.parent_id.id]
                idx = siblings.index(task) + 1
                serial = f"{parent_serial}.{idx}"
            else:
                roots = [t for t in tasks if not t.parent_id]
                idx = roots.index(task) + 1
                serial = str(idx)
            rows.append((task, level, serial, bool(children)))
            for child in children:
                walk(child, level + 1, serial)

        roots = [t for t in tasks if not t.parent_id]
        for root in roots:
            walk(root, 0, "")

        today = datetime.date.today()

        # ── Stats for KPI row ─────────────────────────────────────────
        leaf_tasks = [r for r in rows if not r[3]]
        total_tasks   = len(leaf_tasks)
        completed     = sum(1 for t, l, s, ip in leaf_tasks if (t.progress or 0) == 100)
        in_progress   = sum(1 for t, l, s, ip in leaf_tasks if 0 < (t.progress or 0) < 100)
        not_started   = sum(1 for t, l, s, ip in leaf_tasks if (t.progress or 0) == 0)
        overall_pct   = int(sum((t.progress or 0) for t, l, s, ip in leaf_tasks) / total_tasks) if total_tasks else 0

        # ── Colour palette — Clean Minimal Style ──────────────────────
        C_GREEN_HEADER = "4CAF50"   # bright green header bar
        C_WHITE        = "FFFFFF"
        C_TITLE_FG     = "FFFFFF"

        C_KPI_BG       = "F8F9FA"   # light grey kpi area
        C_KPI_BORDER   = "E0E0E0"

        C_COL_HEADER_BG = "F5F5F5"  # light grey column headers
        C_COL_HEADER_FG = "37474F"
        C_PHASE_BG     = "E8F5E9"   # very light green for phase rows
        C_PHASE_FG     = "2E7D32"   # dark green text for phases
        C_ROW_ODD      = "FFFFFF"
        C_ROW_EVEN     = "FAFAFA"

        # Gantt bar colours (matching legend in screenshot)
        C_BAR_COMPLETE = "4CAF50"   # green  — 100%
        C_BAR_HIGH     = "2196F3"   # blue   — 50–99%
        C_BAR_LOW      = "FF9800"   # orange — 1–49%
        C_BAR_OVERDUE  = "F44336"   # red    — overdue
        C_BAR_ZERO     = "BDBDBD"   # grey   — 0% / not started

        C_WEEKEND_BG   = "F5F5F5"
        C_TODAY_COL    = "F44336"
        C_MILESTONE    = "2E7D32"

        C_GRID         = "E0E0E0"
        C_MONTH_BG     = "ECEFF1"
        C_MONTH_FG     = "37474F"

        def fill(hex_color):
            return PatternFill("solid", fgColor=hex_color)

        def font(bold=False, color="37474F", size=9, italic=False):
            return Font(bold=bold, color=color, size=size, italic=italic,
                        name="Calibri")

        def center(wrap=True):
            return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

        def left(wrap=True):
            return Alignment(horizontal="left", vertical="center", wrap_text=wrap)

        thin = Side(style="thin", color=C_GRID)
        med  = Side(style="medium", color="BDBDBD")
        no   = Side(style=None)

        def border_all():
            return Border(left=thin, right=thin, top=thin, bottom=thin)

        def border_bottom():
            return Border(left=no, right=no, top=no, bottom=thin)

        def border_med_bottom():
            return Border(left=no, right=no, top=no,
                          bottom=Side(style="medium", color="BDBDBD"))

        # ── Fixed columns ─────────────────────────────────────────────
        FIXED_COLS = [
            ("WBS",          7),
            ("TASK / ACTIVITY", 32),
            ("ASSIGNED TO",  16),
            ("PROGRESS",     10),
            ("START",        11),
            ("END",          11),
            ("DAYS",          6),
        ]
        N_FIXED = len(FIXED_COLS)

        # ── Workbook ──────────────────────────────────────────────────
        wb = Workbook()
        ws = wb.active
        ws.title = project.name[:31]
        ws.sheet_view.showGridLines = False

        TOTAL_COLS = N_FIXED + total_days

        # ══════════════════════════════════════════════════════════════
        # ROW 1 — Green title bar
        # ══════════════════════════════════════════════════════════════
        ws.row_dimensions[1].height = 36
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=TOTAL_COLS)
        c = ws.cell(1, 1, value=f"  PROJECT PROGRESS SUMMARY   |   {project.name}")
        c.fill      = fill(C_GREEN_HEADER)
        c.font      = Font(bold=True, color=C_TITLE_FG, size=14, name="Calibri")
        c.alignment = left(wrap=False)

        # ══════════════════════════════════════════════════════════════
        # ROW 2 — KPI Summary bar
        # ══════════════════════════════════════════════════════════════
        ws.row_dimensions[2].height = 40

        kpi_items = [
            ("TOTAL TASKS",  str(total_tasks),  "37474F", "37474F"),
            ("COMPLETED",    str(completed),     "4CAF50", "4CAF50"),
            ("IN PROGRESS",  str(in_progress),   "FF9800", "FF9800"),
            ("NOT STARTED",  str(not_started),   "F44336", "F44336"),
        ]

        kpi_col_width = max(1, N_FIXED // len(kpi_items))

        for ki, (label, value, num_color, _) in enumerate(kpi_items):
            cs = ki * kpi_col_width + 1
            ce = cs + kpi_col_width - 1
            if ki == len(kpi_items) - 1:
                ce = N_FIXED - 2   # leave last 2 cols for overall %
            if ce >= cs:
                ws.merge_cells(start_row=2, start_column=cs, end_row=2, end_column=ce)
            c = ws.cell(2, cs, value=f"{value}  {label}")
            c.fill      = fill(C_KPI_BG)
            c.font      = Font(bold=True, color=num_color, size=11, name="Calibri")
            c.alignment = center()
            c.border    = Border(left=thin, right=thin,
                                  top=Side(style="medium", color=C_GREEN_HEADER),
                                  bottom=thin)

        # Overall % circle (last 2 fixed cols merged)
        ws.merge_cells(start_row=2, start_column=N_FIXED - 1,
                       end_row=2, end_column=N_FIXED)
        c = ws.cell(2, N_FIXED - 1, value=f"{overall_pct}%\nOVERALL")
        c.fill      = fill(C_KPI_BG)
        c.font      = Font(bold=True, color=C_GREEN_HEADER, size=12, name="Calibri")
        c.alignment = center()
        c.border    = Border(left=thin, right=thin,
                              top=Side(style="medium", color=C_GREEN_HEADER),
                              bottom=thin)

        # Gantt area KPI cells (empty, same style)
        for d in range(total_days):
            col = N_FIXED + d + 1
            c = ws.cell(2, col)
            c.fill = fill(C_KPI_BG)
            c.border = Border(top=Side(style="medium", color=C_GREEN_HEADER),
                              bottom=thin)

        # ══════════════════════════════════════════════════════════════
        # ROW 3 — empty spacer
        # ══════════════════════════════════════════════════════════════
        ws.row_dimensions[3].height = 6
        for col in range(1, TOTAL_COLS + 1):
            ws.cell(3, col).fill = fill(C_WHITE)

        # ══════════════════════════════════════════════════════════════
        # ROW 4 — Month labels over Gantt columns
        # ══════════════════════════════════════════════════════════════
        ws.row_dimensions[4].height = 16

        # Fixed col headers (empty for month row)
        for fi in range(N_FIXED):
            c = ws.cell(4, fi + 1)
            c.fill = fill(C_WHITE)

        # Month groups
        month_groups = []
        cur_month = None
        cur_start_d = 0
        for d in range(total_days):
            day_date = range_start + datetime.timedelta(days=d)
            key = (day_date.year, day_date.month)
            if key != cur_month:
                if cur_month:
                    month_groups.append((cur_month, cur_start_d, d - 1))
                cur_month = key
                cur_start_d = d
        month_groups.append((cur_month, cur_start_d, total_days - 1))

        for (yr, mo), dstart, dend in month_groups:
            col_s = N_FIXED + dstart + 1
            col_e = N_FIXED + dend + 1
            if col_e > col_s:
                ws.merge_cells(start_row=4, start_column=col_s,
                               end_row=4, end_column=col_e)
            label = datetime.date(yr, mo, 1).strftime("%b %y").upper()
            c = ws.cell(4, col_s, value=label)
            c.fill      = fill(C_MONTH_BG)
            c.font      = Font(bold=True, color=C_MONTH_FG, size=7, name="Calibri")
            c.alignment = center()
            c.border    = Border(left=med, right=med, top=thin, bottom=thin)
            for d in range(dstart, dend + 1):
                col = N_FIXED + d + 1
                cc = ws.cell(4, col)
                if col != col_s:
                    cc.fill   = fill(C_MONTH_BG)
                    cc.border = Border(left=Side(style="thin", color="CFD8DC"),
                                       right=no, top=thin, bottom=thin)

        # ══════════════════════════════════════════════════════════════
        # ROW 5 — Day-of-week header
        # ══════════════════════════════════════════════════════════════
        ws.row_dimensions[5].height = 16
        DOW = ["M", "T", "W", "T", "F", "S", "S"]

        for fi, (label, width) in enumerate(FIXED_COLS):
            col = fi + 1
            ws.column_dimensions[get_column_letter(col)].width = width
            c = ws.cell(5, col, value=label)
            c.fill      = fill(C_COL_HEADER_BG)
            c.font      = Font(bold=True, color=C_COL_HEADER_FG, size=8, name="Calibri")
            c.alignment = center()
            c.border    = Border(left=thin, right=thin,
                                  top=thin,
                                  bottom=Side(style="medium", color="9E9E9E"))

        for d in range(total_days):
            day_date = range_start + datetime.timedelta(days=d)
            col      = N_FIXED + d + 1
            dow      = DOW[day_date.weekday()]
            ws.column_dimensions[get_column_letter(col)].width = 3.5
            is_today = (day_date == today)
            is_wknd  = day_date.weekday() >= 5
            c = ws.cell(5, col, value=dow)
            if is_today:
                hdr_bg = C_TODAY_COL
                hdr_fg = C_WHITE
            elif is_wknd:
                hdr_bg = "ECEFF1"
                hdr_fg = "90A4AE"
            else:
                hdr_bg = C_COL_HEADER_BG
                hdr_fg = C_COL_HEADER_FG
            c.fill      = fill(hdr_bg)
            c.font      = Font(bold=is_today, color=hdr_fg, size=6, name="Calibri")
            c.alignment = center()
            c.border    = Border(left=thin, right=thin,
                                  top=thin,
                                  bottom=Side(style="medium", color="9E9E9E"))

        # ══════════════════════════════════════════════════════════════
        # ROW 6 — Day number header
        # ══════════════════════════════════════════════════════════════
        ws.row_dimensions[6].height = 14

        for fi in range(N_FIXED):
            c = ws.cell(6, fi + 1)
            c.fill = fill(C_COL_HEADER_BG)
            c.border = border_all()

        for d in range(total_days):
            day_date = range_start + datetime.timedelta(days=d)
            col      = N_FIXED + d + 1
            is_today = (day_date == today)
            is_wknd  = day_date.weekday() >= 5
            c = ws.cell(6, col, value=day_date.day)
            if is_today:
                bg = C_TODAY_COL
                fg = C_WHITE
            elif is_wknd:
                bg = "ECEFF1"
                fg = "90A4AE"
            else:
                bg = C_COL_HEADER_BG
                fg = "78909C"
            c.fill      = fill(bg)
            c.font      = Font(bold=is_today, color=fg, size=6, name="Calibri")
            c.alignment = center()
            c.border    = border_all()

        # ══════════════════════════════════════════════════════════════
        # DATA ROWS — start at row 7
        # ══════════════════════════════════════════════════════════════
        DATA_ROW_START = 7
        child_row_idx = 0  # alternating bg only for leaf rows

        for ri, (task, level, serial, is_parent) in enumerate(rows):
            row_num = DATA_ROW_START + ri
            ws.row_dimensions[row_num].height = 28

            prog      = task.progress or 0
            is_overdue = (task.date_stop and
                          task.date_stop.date() < today and prog < 100)

            # ── Determine background ──────────────────────────────────
            if is_parent:
                bg = C_PHASE_BG
                fg = C_PHASE_FG
                row_font_bold = True
            else:
                bg = C_ROW_ODD if child_row_idx % 2 == 0 else C_ROW_EVEN
                fg = "37474F"
                row_font_bold = False
                child_row_idx += 1

            # ── Progress badge text ───────────────────────────────────
            if is_parent:
                children_tasks = [t for t, l, s, ip in rows
                                   if not ip and s.startswith(serial + ".")]
                if children_tasks:
                    phase_prog = int(sum(t.progress or 0 for t in children_tasks) / len(children_tasks))
                else:
                    phase_prog = 0
                prog_label = f"{phase_prog}%" if phase_prog else ""
            else:
                prog_label = f"{prog:.0f}%"

            # ── Detect milestone (0 days duration) ───────────────────
            duration = 0
            if task.date_start and task.date_stop:
                duration = (task.date_stop.date() - task.date_start.date()).days

            is_milestone = (duration == 0 and not is_parent)

            # ── Fixed column values ───────────────────────────────────
            values = [
                serial,
                task.name or "",
                task.user_id.name if task.user_id else "",
                prog_label,
                task.date_start.strftime("%d %b %y") if task.date_start else "",
                task.date_stop.strftime("%d %b %y")  if task.date_stop  else "",
                str(duration) if task.date_start and task.date_stop else "",
            ]

            for ci, val in enumerate(values):
                col = ci + 1
                c = ws.cell(row_num, col, value=val)
                c.fill = fill(bg)

                if ci == 3 and not is_parent and val:
                    if prog == 100:
                        badge_bg = C_BAR_COMPLETE
                    elif is_overdue:
                        badge_bg = C_BAR_OVERDUE
                    elif prog >= 50:
                        badge_bg = C_BAR_HIGH
                    elif prog > 0:
                        badge_bg = C_BAR_LOW
                    else:
                        badge_bg = C_BAR_ZERO
                    c.fill = fill(badge_bg)
                    c.font = Font(bold=True, color=C_WHITE, size=8, name="Calibri")
                else:
                    c.font = Font(bold=row_font_bold, color=fg, size=9, name="Calibri")

                c.alignment = left() if ci == 1 else center()
                c.border    = Border(left=thin, right=thin,
                                      top=no, bottom=thin)

            # ── Gantt bar ─────────────────────────────────────────────
            t_start = task.date_start.date() if task.date_start else None
            t_end   = task.date_stop.date()  if task.date_stop  else None

            if prog == 100:
                bar_color = C_BAR_COMPLETE
            elif is_overdue:
                bar_color = C_BAR_OVERDUE
            elif is_parent:
                bar_color = "A5D6A7"   # lighter green for phase bars
            elif prog >= 50:
                bar_color = C_BAR_HIGH
            elif prog > 0:
                bar_color = C_BAR_LOW
            else:
                bar_color = C_BAR_ZERO

            # ── Gantt bar — merged pill style ─────────────────────────
            DAY_SEP = Side(style="thin", color="CFD8DC")

            # First pass: fill all non-bar cells with grid line
            for d in range(total_days):
                day_date = range_start + datetime.timedelta(days=d)
                col      = N_FIXED + d + 1
                c        = ws.cell(row_num, col)
                is_wknd  = day_date.weekday() >= 5

                is_in_bar = (t_start and t_end and
                             t_start <= day_date <= t_end and
                             not is_milestone)

                if is_milestone and t_start and day_date == t_start:
                    c.value     = "◆"
                    c.fill      = fill(bg)
                    c.font      = Font(bold=True, color=C_MILESTONE, size=10, name="Calibri")
                    c.alignment = center()
                    c.border    = Border(top=no, bottom=thin, left=DAY_SEP, right=no)

                elif not is_in_bar:
                    cell_bg = C_WEEKEND_BG if is_wknd else bg
                    c.fill   = fill(cell_bg)
                    c.border = Border(top=no, bottom=thin, left=DAY_SEP, right=no)

            # Second pass: draw merged bar
            if t_start and t_end and not is_milestone:
                bar_hex = bar_color
                bg_hex  = bg

                bar_col_start = N_FIXED + (t_start - range_start).days + 1
                bar_col_end   = N_FIXED + (t_end   - range_start).days + 1
                bar_col_start = max(bar_col_start, N_FIXED + 1)
                bar_col_end   = min(bar_col_end,   N_FIXED + total_days)

                if bar_col_start <= bar_col_end:
                    if bar_col_end > bar_col_start:
                        ws.merge_cells(
                            start_row=row_num, start_column=bar_col_start,
                            end_row=row_num,   end_column=bar_col_end
                        )

                    c = ws.cell(row_num, bar_col_start)
                    c.fill = GradientFill(
                        type="linear",
                        degree=90,
                        stop=(
                            Stop(position=0.0,  color=bg_hex),
                            Stop(position=0.28, color=bg_hex),
                            Stop(position=0.33, color=bar_hex),
                            Stop(position=0.67, color=bar_hex),
                            Stop(position=0.72, color=bg_hex),
                            Stop(position=1.0,  color=bg_hex),
                        )
                    )
                    c.alignment = center(wrap=False)
                    grid_line = Side(style="thin",   color="E0E0E0")
                    bg_pad    = Side(style="medium",  color=bg_hex)
                    no_side   = Side(style=None)
                    # Step 1: stamp bottom grid line on every cell in the span
                    # (merged-cell borders only apply at the anchor in openpyxl,
                    #  so we must touch each cell individually for the bottom line)
                    for _d in range(bar_col_start, bar_col_end + 1):
                        _c = ws.cell(row_num, _d)
                        _c.border = Border(
                            top    = no_side,
                            bottom = grid_line,
                            left   = no_side,
                            right  = no_side,
                        )
                    # Step 2: anchor cell gets the pill padding on top/left/right
                    c.border = Border(
                        top    = bg_pad,
                        bottom = grid_line,
                        left   = bg_pad,
                        right  = bg_pad,
                    )

            # ── Today marker line ─────────────────────────────────────
            if range_start <= today <= range_end:
                today_offset = (today - range_start).days
                today_col    = N_FIXED + today_offset + 1
                c = ws.cell(row_num, today_col)
                c.border = Border(left=Side(style="medium", color=C_TODAY_COL),
                                   right=no, top=no, bottom=thin)

        # ══════════════════════════════════════════════════════════════
        # LEGEND ROW
        # ══════════════════════════════════════════════════════════════
        legend_row = DATA_ROW_START + len(rows) + 1
        ws.row_dimensions[legend_row] = ws.row_dimensions.get(legend_row) or \
                                         ws.row_dimensions[legend_row]
        ws.row_dimensions[legend_row].height = 20

        legend_items = [
            ("■ 100% Completed",        C_BAR_COMPLETE),
            ("■ In Progress (50% - 99%)", C_BAR_HIGH),
            ("■ In Progress (1% - 49%)", C_BAR_LOW),
            ("■ Overdue",               C_BAR_OVERDUE),
            ("■ Not Started",           C_BAR_ZERO),
            ("◆ Milestone",             C_MILESTONE),
        ]

        leg_col = 1
        for label, color in legend_items:
            ws.merge_cells(start_row=legend_row, start_column=leg_col,
                           end_row=legend_row, end_column=leg_col + 1)
            c = ws.cell(legend_row, leg_col, value=label)
            c.fill      = fill(C_WHITE)
            c.font      = Font(bold=False, color=color, size=8, name="Calibri")
            c.alignment = left(wrap=False)
            leg_col += 2

        # ── Freeze panes ──────────────────────────────────────────────
        ws.freeze_panes = ws.cell(DATA_ROW_START, N_FIXED + 1)

        # ── Auto-filter ───────────────────────────────────────────────
        ws.auto_filter.ref = (
            f"A5:{get_column_letter(N_FIXED)}{DATA_ROW_START + len(rows) - 1}"
        )

        # ── Save ──────────────────────────────────────────────────────
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = f"{project.name}_Gantt_{today.strftime('%Y-%m-%d')}.xlsx"
        return request.make_response(
            buf.getvalue(),
            headers=[
                ('Content-Type',
                 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', content_disposition(filename)),
            ]
        )